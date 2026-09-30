from openpyxl import load_workbook

from slug import slug

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}
ENCABEZADO = "indicador"
FILAS = (("poblacion", "poblacion"), ("eventos", "eventos"), ("dias", "dias"))
MAX_FILAS_ENCABEZADO = 20


def _slug_de(celda):
    return slug(celda.value) if isinstance(celda.value, str) else ""


def _encabezado(ws):
    for fila in range(1, min(ws.max_row, MAX_FILAS_ENCABEZADO) + 1):
        for celda in ws[fila]:
            if _slug_de(celda) != ENCABEZADO:
                continue
            columnas = {c.column: MESES[_slug_de(c)] for c in ws[fila] if c.column > celda.column and _slug_de(c) in MESES}
            if columnas:
                return fila, celda.column, columnas
    return None


def _titulo(ws, fila):
    for f in range(1, fila):
        for celda in ws[f]:
            if isinstance(celda.value, str) and celda.value.strip():
                return slug(celda.value)
    return ""


def _filas_de_etiquetas(ws, fila, col):
    filas = {}
    for f in range(fila + 1, ws.max_row + 1):
        etiqueta = _slug_de(ws.cell(row=f, column=col))
        for clave, prefijo in FILAS:
            if clave not in filas and etiqueta.startswith(prefijo):
                filas[clave] = f
    return filas if len(filas) == len(FILAS) else None


def _entero(v):
    return int(round(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def leer_hoja(ws):
    encabezado = _encabezado(ws)
    if encabezado is None:
        return None
    fila, col, columnas = encabezado
    filas = _filas_de_etiquetas(ws, fila, col)
    if filas is None:
        return None
    meses = {
        mes: {clave: _entero(ws.cell(row=filas[clave], column=c).value) for clave, _ in FILAS}
        for c, mes in columnas.items()
    }
    return {"titulo": _titulo(ws, fila), "meses": meses}


def leer_libro(ruta):
    wb = load_workbook(ruta, data_only=True)
    return [(ws.title, leer_hoja(ws)) for ws in wb]
