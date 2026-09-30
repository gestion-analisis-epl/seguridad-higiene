from openpyxl import Workbook

from leer_indices import MESES, leer_hoja, leer_libro

NOMBRES = list(MESES)


def hoja_indices(ws, titulo="Indices de Ciudad A", fila=3, col=2, datos=None, etiquetas=("Población", "HHT", "Eventos", "Días", "IF")):
    ws.cell(row=1, column=col, value=titulo)
    ws.cell(row=fila, column=col, value="Indicador")
    for i, mes in enumerate(NOMBRES, start=1):
        ws.cell(row=fila, column=col + i, value=mes.capitalize())
    ws.cell(row=fila, column=col + 13, value="Anual")
    datos = datos or {}
    for n, etiqueta in enumerate(etiquetas, start=1):
        ws.cell(row=fila + n, column=col, value=etiqueta)
        for mes, valores in datos.items():
            if n - 1 < len(valores):
                ws.cell(row=fila + n, column=col + mes, value=valores[n - 1])
    return ws


def test_detecta_el_layout_por_texto_en_cualquier_posicion():
    wb = Workbook()
    ws = hoja_indices(wb.active, fila=6, col=4, datos={1: [10, "=B1", 2, 7]})
    datos = leer_hoja(ws)
    assert datos["titulo"] == "indices-de-ciudad-a"
    assert datos["meses"][1] == {"poblacion": 10, "eventos": 2, "dias": 7}


def test_mapea_los_doce_meses_y_acepta_acentos_y_mayusculas():
    assert MESES["enero"] == 1 and MESES["diciembre"] == 12 and len(MESES) == 12
    wb = Workbook()
    ws = hoja_indices(wb.active, datos={m: [m, None, 0, 0] for m in range(1, 13)})
    ws.cell(row=3, column=4, value="FEBRERO")
    assert sorted(leer_hoja(ws)["meses"]) == list(range(1, 13))
    assert leer_hoja(ws)["meses"][12]["poblacion"] == 12


def test_ignora_filas_calculadas_y_valores_no_numericos():
    wb = Workbook()
    ws = hoja_indices(wb.active, datos={1: ["x", None, 1.0, None]})
    assert leer_hoja(ws)["meses"][1] == {"poblacion": 0, "eventos": 1, "dias": 0}


def test_hoja_sin_layout_es_formato_no_reconocido():
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Otra cosa")
    ws.cell(row=2, column=1, value="Indicador")
    assert leer_hoja(ws) is None


def test_sin_fila_de_dias_no_se_reconoce():
    wb = Workbook()
    ws = hoja_indices(wb.active, etiquetas=("Población", "Eventos"))
    assert leer_hoja(ws) is None


def test_leer_libro_devuelve_todas_las_hojas(tmp_path):
    wb = Workbook()
    hoja_indices(wb.active).title = "A"
    wb.create_sheet("Extra").cell(row=1, column=1, value="nada")
    ruta = tmp_path / "i.xlsx"
    wb.save(ruta)
    hojas = leer_libro(ruta)
    assert [t for t, _ in hojas] == ["A", "Extra"]
    assert hojas[0][1] is not None and hojas[1][1] is None
