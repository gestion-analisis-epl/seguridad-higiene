from openpyxl import Workbook

from leer_excel import OFF, encontrar_layout, es_hoja_de_colaboradores, filas_de_datos, offsets_de, segmentar_bloques


def poner(ws, fila, base, off, valor):
    ws.cell(row=fila, column=base + off + 1, value=valor)


def hoja_completa(fila_nombre=5, base=5):
    wb = Workbook()
    ws = wb.active
    poner(ws, fila_nombre, base, 0, "Nombre")
    poner(ws, fila_nombre - 1, base, OFF["botas"], "Botas")
    poner(ws, fila_nombre - 1, base, OFF["guantes"], "Guantes")
    poner(ws, fila_nombre - 2, base, OFF["acc_fecha"], "Accidentes")
    return ws


def test_encuentra_el_layout_completo():
    assert encontrar_layout(hoja_completa(5, 5)) == (5, 5)
    assert encontrar_layout(hoja_completa(4, 6)) == (4, 6)


def test_rechaza_hojas_reducidas():
    wb = Workbook()
    ws = wb.active
    ws.cell(row=3, column=6, value="Personal")
    ws.cell(row=3, column=9, value="Curso")
    assert encontrar_layout(ws) is None


def test_lee_filas_despues_del_encabezado():
    ws = hoja_completa(5, 5)
    poner(ws, 6, 5, 0, "Ana")
    poner(ws, 7, 5, 0, "Beto")
    filas = filas_de_datos(ws, (5, 5))
    assert [n for n, _ in filas][:2] == [6, 7]


def hoja_con_ancla(ancla, desplazo=0, fila_nombre=4, base=5):
    wb = Workbook()
    ws = wb.active
    poner(ws, fila_nombre, base, 0, ancla)
    poner(ws, fila_nombre - 1, base, OFF["botas"] + desplazo, "Botas")
    poner(ws, fila_nombre - 1, base, OFF["guantes"] + desplazo, "Guantes")
    poner(ws, fila_nombre - 2, base, OFF["acc_fecha"] + desplazo, "Accidentes")
    return ws


def test_acepta_personal_como_ancla_si_el_bloque_valida():
    assert encontrar_layout(hoja_con_ancla("Personal")) == (4, 5)


def test_personal_sin_bloques_no_es_layout():
    wb = Workbook()
    ws = wb.active
    poner(ws, 4, 5, 0, "Personal")
    assert encontrar_layout(ws) is None


def test_prefiere_nombre_sobre_personal():
    ws = hoja_con_ancla("Personal", fila_nombre=6, base=2)
    poner(ws, 4, 8, 0, "Nombre")
    poner(ws, 3, 8, OFF["botas"], "Botas")
    poner(ws, 3, 8, OFF["guantes"], "Guantes")
    poner(ws, 2, 8, OFF["acc_fecha"], "Accidentes")
    assert encontrar_layout(ws) == (4, 8)


def test_columna_extra_desplaza_los_bloques():
    ws = hoja_con_ancla("Personal", desplazo=1)
    layout = encontrar_layout(ws)
    assert layout == (4, 5)
    off = offsets_de(ws, layout)
    for clave in ("norma", "cap_fecha", "cap_venc", "botas", "guantes", "acc_fecha", "acc_trayecto", "acc_laboral"):
        assert off[clave] == OFF[clave] + 1
    for clave in ("linea", "ciudad", "ingreso", "cuadrilla"):
        assert off[clave] == OFF[clave]


def test_bloques_por_encabezado_repetido():
    filas = [(6, ["x", "Ana"]), (7, ["x", "Beto"]), (8, ["x", "Personal"]), (9, ["x", None]),
             (10, ["x", "Nombre"]), (11, ["x", "Ana"])]
    bloques = segmentar_bloques(filas, 1)
    assert [[n for n, _ in b if n in (6, 7, 11)] for b in bloques] == [[6, 7], [11]]


def test_linea_y_ciudad_se_ubican_por_su_encabezado():
    ws = hoja_con_ancla("Personal")
    poner(ws, 2, 5, -5, "Linea de negocio")
    poner(ws, 2, 5, -2, "Ciudad")
    off = offsets_de(ws, (4, 5))
    assert off["linea"] == -5 and off["ciudad"] == -2


def test_sin_encabezados_de_linea_y_ciudad_usa_los_predeterminados():
    off = offsets_de(hoja_con_ancla("Personal"), (4, 5))
    assert off["linea"] == OFF["linea"] and off["ciudad"] == OFF["ciudad"]


def test_sin_exclusiones_toda_hoja_es_candidata():
    assert es_hoja_de_colaboradores("Hoja Uno")
    assert es_hoja_de_colaboradores("Hoja Uno", [])


def test_excluye_por_subcadena_sobre_el_slug_del_titulo():
    assert not es_hoja_de_colaboradores("Resumen Área Dos", ["area dos"])
    assert es_hoja_de_colaboradores("Hoja Uno", ["area dos"])
