from datetime import datetime

from openpyxl import Workbook

from leer_excel import OFF
from migrar import leer_todo

BASE = 5


def poner_hoja(ws, filas):
    ws.cell(row=4, column=BASE + 1, value="Nombre")
    ws.cell(row=3, column=BASE + OFF["botas"] + 1, value="Botas")
    ws.cell(row=3, column=BASE + OFF["guantes"] + 1, value="Guantes")
    ws.cell(row=2, column=BASE + OFF["acc_fecha"] + 1, value="Accidentes")
    for i, (nombre, ciudad, linea) in enumerate(filas, start=5):
        for off, v in ((0, nombre), (OFF["ciudad"], ciudad), (OFF["linea"], linea)):
            ws.cell(row=i, column=BASE + off + 1, value=v)


def libro(tmp_path, filas, otras=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja Colaboradores"
    poner_hoja(ws, filas)
    for titulo in otras:
        poner_hoja(wb.create_sheet(titulo), [("Eva", "Ciudad C", "LINEA Z")])
    ruta = tmp_path / "t.xlsx"
    wb.save(ruta)
    return ruta


def test_filas_sin_ciudad_o_linea_no_llegan_ni_a_los_catalogos(tmp_path):
    ruta = libro(tmp_path, [("Ana", "Ciudad A", "LINEA X"), ("Beto", None, "LINEA X"), ("Cris", "Ciudad A", None)])
    registros, reporte, catalogos = leer_todo(ruta, datetime(2026, 9, 30))
    assert len(registros) == 1
    assert all(i["valor"] for items in catalogos.values() for i in items)
    assert sum("sin ciudad" in a for a in reporte[0]["avisos"]) == 1
    assert sum("sin linea de negocio" in a for a in reporte[0]["avisos"]) == 1


def test_reporte_cuenta_uniformes_sin_fecha(tmp_path):
    ruta = libro(tmp_path, [("Ana", "Ciudad A", "LINEA X")])
    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    wb.active.cell(row=5, column=BASE + OFF["botas"] + 1, value=26)
    wb.save(ruta)
    _, reporte, _ = leer_todo(ruta, datetime(2026, 9, 30))
    assert reporte[0]["uniformes_sin_fecha"] == 1


def test_catalogos_toman_la_etiqueta_del_texto_original_y_gana_el_primero(tmp_path):
    ruta = libro(tmp_path, [("Ana", "Ciudad B - Área Uno", "Linea X"), ("Beto", "CIUDAD B-Area uno", "LINEA X")])
    _, _, catalogos = leer_todo(ruta, datetime(2026, 9, 30))
    assert catalogos["ciudades"] == [{"valor": "ciudad-b", "etiqueta": "Ciudad B"}]
    assert catalogos["areas"] == [{"valor": "area-uno", "etiqueta": "Área Uno"}]
    assert catalogos["lineas_negocio"] == [{"valor": "linea-x", "etiqueta": "Linea X"}]


def test_la_config_excluye_hojas_y_desactiva_la_division(tmp_path):
    ruta = libro(tmp_path, [("Ana", "Ciudad B-Area Uno", "LINEA X")], otras=["Resumen Extra"])
    registros, reporte, _ = leer_todo(ruta, datetime(2026, 9, 30))
    assert len(registros) == 2 and len(reporte) == 2
    config = {"separador_area": None, "excluir_hojas": ["resumen"]}
    registros, reporte, _ = leer_todo(ruta, datetime(2026, 9, 30), config)
    assert [h["hoja"] for h in reporte] == ["Hoja Colaboradores"]
    assert registros[0]["colaborador"]["ciudad"] == "ciudad-b-area-uno"


def test_simular_y_commit_son_excluyentes_y_no_conectan(monkeypatch):
    import sys

    import pytest

    import migrar
    monkeypatch.setattr(sys, "argv", ["migrar.py", "--excel", "x.xlsx", "--simular", "--commit"])
    with pytest.raises(SystemExit) as e:
        migrar.main()
    assert e.value.code != 0


def test_simular_sin_destino_falla_antes_de_leer_o_conectar(monkeypatch):
    import sys

    import pytest

    import migrar
    monkeypatch.delenv("MIGRACION_PROYECTO", raising=False)
    monkeypatch.delenv("MIGRACION_BASE", raising=False)
    monkeypatch.setattr(sys, "argv", ["migrar.py", "--excel", "no-existe.xlsx", "--simular"])
    with pytest.raises(SystemExit) as e:
        migrar.main()
    assert "--simular" in str(e.value)
